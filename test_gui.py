import customtkinter as ctk
from tkinter import filedialog
import tkinter as tk
import os
 
import docx
import PyPDF2
import openpyxl
 
# Mock GPT4All class for demonstration; replace with your actual instance
class GPT4All:
    def generate(self, prompt, **kwargs):
        return ["This is a summarized version of the file content."]  # Replace with real generation
 
# Your GPT-based response function
def generate_response(gpt4all_instance, first_name: str, last_name: str, subject: str, body: str, rag_context: str = "", additional_prompt: str = "") -> str:
    prompt = (
        "You are a professional, empathetic cruise line guest relations specialist"
        "You respond clearly, politely, and with care, using advanced English at a C1 proficiency level.\n\n"
        "If the guest's concern involves calculation, refund, or policies, think through the issue step by step, "
        "then summarize the final answer in a polished reply.\n\n"
        f"Guest Name: {first_name} {last_name}\n"
        f"Subject: {subject}\n"
        f"Body of the Guest's Email:\n{body}\n"
    )
 
    if additional_prompt:
        prompt += f"\nImportant background about this case:\n{additional_prompt.strip()}\n"
 
    if rag_context:
        prompt += f"\nBackground documents relevant to this case:\n{rag_context[:3000]}\n"
 
    prompt += (
        "\nCompose a well-structured C1-level English reply addressing the guest by their first name."
        "Always sound professional, empathetic, and solution-oriented. Avoid generic phrasing. Close with reassurance or gratitude.\n"
        "If the guest's request cannot be fully granted, avoid saying 'we cannot' directly. "
        "Instead, suggest alternative options in a positive way.\n"
        "Always focus on what *can* be done.\n"
        "Reply:\n"
    )
 
    response = gpt4all_instance.generate(prompt, max_tokens=512, temp=0.7, top_k=40, top_p=0.9, repeat_penalty=1.1, streaming=False)
    return "".join(response) if hasattr(response, '__iter__') else str(response)
 
class EmailResponderGUI(ctk.CTk):
    def __init__(self):
        super().__init__()
        self.title("AI Email Generator")
        self.geometry("900x600")
        ctk.set_appearance_mode("dark")
        ctk.set_default_color_theme("dark-blue")
 
        self.gpt = GPT4All()
        self.attached_files = []
        self.database_file = "database.txt"
 
        self.tabview = ctk.CTkTabview(self)
        self.tabview.pack(fill="both", expand=True, padx=20, pady=20)
 
        self.input_tab = self.tabview.add("Guest Email Input")
        self.output_tab = self.tabview.add("AI Response Output")
 
        self.create_input_tab()
        self.create_output_tab()
 
    def create_input_tab(self):
        self.first_name = ctk.CTkEntry(self.input_tab, placeholder_text="First Name")
        self.first_name.grid(row=0, column=0, padx=10, pady=10, sticky="ew")
 
        self.last_name = ctk.CTkEntry(self.input_tab, placeholder_text="Last Name")
        self.last_name.grid(row=0, column=1, padx=10, pady=10, sticky="ew")
 
        self.subject = ctk.CTkEntry(self.input_tab, placeholder_text="Email Subject")
        self.subject.grid(row=1, column=0, columnspan=2, padx=10, pady=10, sticky="ew")
 
        self.body = ctk.CTkTextbox(self.input_tab, height=150)
        self.body.grid(row=2, column=0, columnspan=2, padx=10, pady=10, sticky="nsew")
        self.body.insert("0.0", "Body of the email...")
 
        self.concern = ctk.CTkEntry(self.input_tab, placeholder_text="Concern")
        self.concern.grid(row=3, column=0, columnspan=2, padx=10, pady=10, sticky="ew")
 
        self.tone_combo = ctk.CTkComboBox(self.input_tab, values=["Formal", "Friendly", "Empathetic", "Apologetic"])
        self.tone_combo.set("Select Tone")
        self.tone_combo.grid(row=4, column=0, columnspan=2, padx=10, pady=10)
 
        self.attach_button = ctk.CTkButton(self.input_tab, text="Attach Files", command=self.attach_files)
        self.attach_button.grid(row=5, column=0, padx=10, pady=10)
 
        self.generate_button = ctk.CTkButton(self.input_tab, text="Generate", command=self.generate_response)
        self.generate_button.grid(row=5, column=1, padx=10, pady=10)
 
        self.regenerate_button = ctk.CTkButton(self.input_tab, text="Regenerate", command=self.generate_response)
        self.regenerate_button.grid(row=6, column=0, columnspan=2, padx=10, pady=10)
 
        self.attach_label = ctk.CTkLabel(self.input_tab, text="No files attached.")
        self.attach_label.grid(row=7, column=0, columnspan=2, padx=10, pady=10)
 
        self.input_tab.grid_rowconfigure(2, weight=1)
        self.input_tab.grid_columnconfigure((0, 1), weight=1)
 
    def create_output_tab(self):
        self.generated_subject = ctk.CTkEntry(self.output_tab, placeholder_text="AI Reply Subject")
        self.generated_subject.pack(padx=10, pady=10, fill="x")
 
        self.generated_body = ctk.CTkTextbox(self.output_tab)
        self.generated_body.pack(padx=10, pady=10, fill="both", expand=True)
 
    def attach_files(self):
        file_paths = filedialog.askopenfilenames(
            title="Select Reference Documents",
            filetypes=[("All Files", "*.*"),
                       ("Text Files", "*.txt"),
                       ("PDF Files", "*.pdf"),
                       ("Word Documents", "*.docx"),
                       ("Excel Files", "*.xlsx"),
                       ("Images", "*.jpg *.jpeg *.png")]
        )
        self.attached_files = file_paths
        if file_paths:
            self.attach_label.configure(text=f"{len(file_paths)} files attached.")
 
        concern_tag = self.concern.get().strip()
        if not concern_tag:
            self.attach_label.configure(text="Add a concern before attaching.")
            return
 
        for file in file_paths:
            content = self.extract_file_text(file)
            if content:
                summary = generate_response(self.gpt, "System", "Summary", "Reference", content)
                with open(self.database_file, "a", encoding="utf-8") as f:
                    f.write(f"[{concern_tag}]\n{summary.strip()}\n\n")
 
    def extract_file_text(self, filepath):
        try:
            if filepath.endswith(".txt"):
                with open(filepath, "r", encoding="utf-8") as f:
                    return f.read()
            elif filepath.endswith(".pdf"):
                with open(filepath, "rb") as f:
                    reader = PyPDF2.PdfReader(f)
                    return "\n".join([page.extract_text() for page in reader.pages if page.extract_text()])
            elif filepath.endswith(".docx"):
                doc = docx.Document(filepath)
                return "\n".join([para.text for para in doc.paragraphs])
            elif filepath.endswith(".xlsx"):
                wb = openpyxl.load_workbook(filepath)
                text = []
                for sheet in wb:
                    for row in sheet.iter_rows(values_only=True):
                        text.append(" | ".join([str(cell) for cell in row if cell]))
                return "\n".join(text)
            return ""
        except Exception as e:
            print(f"Error reading {filepath}: {e}")
            return ""
 
    def find_rag_context(self, concern_tag):
        try:
            with open(self.database_file, "r", encoding="utf-8") as f:
                content = f.read()
            sections = content.split("\n[")
            for sec in sections:
                if sec.startswith(concern_tag + "]"):
                    return sec.split("]", 1)[-1].strip()
        except FileNotFoundError:
            pass
        return ""
 
    def generate_response(self):
        first_name = self.first_name.get()
        last_name = self.last_name.get()
        subject = self.subject.get()
        tone = self.tone_combo.get()
        body = self.body.get("1.0", "end").strip()
        concern_tag = self.concern.get().strip()
 
        rag_context = self.find_rag_context(concern_tag)
 
        reply = generate_response(
            self.gpt, first_name, last_name, subject, body, rag_context=rag_context
        )
 
        self.generated_subject.delete(0, "end")
        self.generated_subject.insert(0, f"Re: {subject}")
 
        self.generated_body.delete("1.0", "end")
        self.generated_body.insert("1.0", reply)
 
 
if __name__ == "__main__":
    app = EmailResponderGUI()
    app.mainloop()